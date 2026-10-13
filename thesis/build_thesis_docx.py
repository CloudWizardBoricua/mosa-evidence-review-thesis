"""Build the separately revised thesis from inspectable JSON source.

When Codex uses this builder, its document artifact marker precedes creation.
Rendering and review of every page remain required before final delivery.
"""
import argparse
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

BASE = Path(__file__).resolve().parent
ROOT = BASE.parent if (BASE.parent / 'mkdocs.yml').exists() else BASE.parent / 'mosa-evidence-review-thesis'
COLORS = {
    'blue': ('205493', 'EAF1F8'),
    'purple': ('67459B', 'F0ECF6'),
    'green': ('286B43', 'EDF5EF'),
    'amber': ('9B6800', 'FFF6DF'),
    'crimson': ('841617', 'F8EBEB'),
}
FIGURE_ALIASES = {
    'system-topology': ['system_topology', 'system-topology', 'topology'],
    'review-sequence': ['review_sequence', 'review-sequence'],
    'approval-rollback': ['approval_rollback', 'approval-rollback'],
    'benchmark-split': ['benchmark_split', 'benchmark-split'],
    'accuracy-chart': ['accuracy', 'accuracy_chart', 'accuracy-chart'],
    'effort-latency-cost': ['effort_latency_cost', 'effort-latency-cost', 'performance'],
    'migration': ['migration'],
}


def load_json(path, default=None):
    return json.loads(Path(path).read_text(encoding='utf-8')) if path else (default if default is not None else {})


def apply_substitutions(value, replacements):
    if isinstance(value, str):
        for key, replacement in replacements.items():
            token = key if key.startswith('{{') else '{{' + key + '}}'
            value = value.replace(token, str(replacement))
        return value
    if isinstance(value, dict):
        return {key: apply_substitutions(child, replacements) for key, child in value.items()}
    if isinstance(value, list):
        return [apply_substitutions(child, replacements) for child in value]
    return value


def iter_strings(value):
    if isinstance(value, str):
        yield value
    elif isinstance(value, dict):
        for child in value.values():
            yield from iter_strings(child)
    elif isinstance(value, list):
        for child in value:
            yield from iter_strings(child)


def font_run(run, size=12, bold=None):
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.insert(0, fonts)
    for attr in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn('w:' + attr), 'Times New Roman')
    return run


def setup_styles(doc):
    for style in doc.styles:
        if style.type == 1:
            style.font.name = 'Times New Roman'
            style.font.size = Pt(12)
            style.font.color.rgb = RGBColor(0, 0, 0)
    normal = doc.styles['Normal']
    normal.paragraph_format.line_spacing = 2.0
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Inches(.5)
    normal.paragraph_format.widow_control = True
    for name in ('Title', 'Heading 1', 'Heading 2', 'Heading 3'):
        style = doc.styles[name]
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(12)
        style.font.bold = name != 'Title'
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.first_line_indent = Inches(0)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    for name in ('Caption', 'TOC Heading'):
        style = doc.styles[name]
        style.font.size = Pt(12)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.first_line_indent = Inches(0)
        style.paragraph_format.space_before = Pt(6)
        style.paragraph_format.space_after = Pt(6)
    for name in ('Front Heading', 'Guide Text', 'Reference Text', 'Glossary Text', 'TOC Entry'):
        if name not in doc.styles:
            doc.styles.add_style(name, 1)
        style = doc.styles[name]
        style.font.name = 'Times New Roman'
        style.font.size = Pt(12)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.first_line_indent = Inches(0)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.widow_control = True
    doc.styles['Front Heading'].font.bold = True
    doc.styles['Front Heading'].paragraph_format.space_after = Pt(14)
    doc.styles['Reference Text'].paragraph_format.line_spacing = 2.0
    doc.styles['Reference Text'].paragraph_format.left_indent = Inches(.5)
    doc.styles['Reference Text'].paragraph_format.first_line_indent = Inches(-.5)
    doc.styles['Reference Text'].paragraph_format.space_after = Pt(0)
    doc.styles['Glossary Text'].paragraph_format.space_after = Pt(6)
    doc.styles['TOC Entry'].paragraph_format.tab_stops.add_tab_stop(Inches(6.25), 2, 1)
    # Remove possible ornamental rules inherited from Word themes.
    for name in ('Title', 'Heading 1', 'Heading 2', 'Front Heading'):
        ppr = doc.styles[name].element.find(qn('w:pPr'))
        if ppr is not None:
            for element in list(ppr.findall(qn('w:pBdr'))):
                ppr.remove(element)


def configure_section(section, number_format=None, start=None):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(.5)
    section.footer_distance = Inches(.5)
    section.footer.is_linked_to_previous = False
    section.header.is_linked_to_previous = False
    for p in section.footer.paragraphs:
        p.clear()
    page_settings = section._sectPr.find(qn('w:pgNumType'))
    if page_settings is not None:
        section._sectPr.remove(page_settings)
    if number_format:
        page_settings = OxmlElement('w:pgNumType')
        page_settings.set(qn('w:fmt'), number_format)
        page_settings.set(qn('w:start'), str(start))
        section._sectPr.append(page_settings)
        p = section.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.line_spacing = 1.0
        add_field(p, 'PAGE', 'iv' if number_format == 'lowerRoman' else '1')


def add_field(paragraph, instruction, cached=''):
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    instruction_element = OxmlElement('w:instrText')
    instruction_element.set(qn('xml:space'), 'preserve')
    instruction_element.text = ' ' + instruction + ' '
    separate = OxmlElement('w:fldChar')
    separate.set(qn('w:fldCharType'), 'separate')
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    font_run(paragraph.add_run())._r.append(begin)
    font_run(paragraph.add_run())._r.append(instruction_element)
    font_run(paragraph.add_run())._r.append(separate)
    font_run(paragraph.add_run(cached))
    font_run(paragraph.add_run())._r.append(end)


def add_external_link(paragraph, text, url):
    relationship = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), relationship)
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    fonts = OxmlElement('w:rFonts')
    fonts.set(qn('w:ascii'), 'Times New Roman')
    fonts.set(qn('w:hAnsi'), 'Times New Roman')
    props.append(fonts)
    size = OxmlElement('w:sz')
    size.set(qn('w:val'), '24')
    props.append(size)
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '205493')
    props.append(color)
    underline = OxmlElement('w:u')
    underline.set(qn('w:val'), 'single')
    props.append(underline)
    run.append(props)
    content = OxmlElement('w:t')
    content.text = text
    run.append(content)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


class Builder:
    def __init__(self, args):
        self.args = args
        self.content = apply_substitutions(load_json(args.content), load_json(args.substitutions))
        self.doc = Document()
        setup_styles(self.doc)
        configure_section(self.doc.sections[0])
        self.page_map = load_json(args.page_map)
        self.figure_map = load_json(args.figure_map)
        self.bookmark_id = 1
        self.figure_count = 0
        self.table_count = 0
        self.figures = []
        self.missing_assets = []
        self.bookmarks = []
        self.original_placeholders = sorted(set(re.findall(r'\{\{study[^}]*\}\}', '\n'.join(iter_strings(self.content)))))
        if self.original_placeholders and not args.allow_placeholders:
            raise ValueError('Unresolved empirical placeholders: ' + ', '.join(self.original_placeholders))
        if args.allow_placeholders:
            self.content = apply_substitutions(self.content, {p: 'Evaluation result pending. No empirical claim is supplied for this measure.' for p in self.original_placeholders})
        self.doc.core_properties.title = self.content['metadata']['title']
        self.doc.core_properties.author = self.content['metadata']['author']
        self.doc.core_properties.subject = 'Condensed thesis revision with bounded agent supported evidence review'
        self.doc.core_properties.keywords = '; '.join(self.content['front']['keywords'])
        self.doc.core_properties.comments = 'Original completed thesis remains unchanged. This revision requires author and committee review.'
        for chapter in self.content['chapters']:
            for sec in chapter['sections']:
                if sec.get('figure'):
                    self.figures.append(sec['figure'])
        for appendix in self.content.get('appendices', []):
            for sec in appendix['sections']:
                if sec.get('figure'):
                    self.figures.append(sec['figure'])

    def bookmark(self, paragraph, name):
        name = re.sub('[^A-Za-z0-9_]', '_', name)[:39]
        start = OxmlElement('w:bookmarkStart')
        start.set(qn('w:id'), str(self.bookmark_id))
        start.set(qn('w:name'), name)
        end = OxmlElement('w:bookmarkEnd')
        end.set(qn('w:id'), str(self.bookmark_id))
        # Paragraph properties precede inline content in the OOXML schema.
        paragraph._p.insert(1 if paragraph._p.find(qn('w:pPr')) is not None else 0, start)
        paragraph._p.append(end)
        self.bookmark_id += 1
        self.bookmarks.append(name)
        return name

    def para(self, text='', style=None, center=False, space_before=None, space_after=None, bold=None):
        p = self.doc.add_paragraph(style=style)
        font_run(p.add_run(text), bold=bold)
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Inches(0)
        if space_before is not None:
            p.paragraph_format.space_before = Pt(space_before)
        if space_after is not None:
            p.paragraph_format.space_after = Pt(space_after)
        return p

    def break_page(self):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        font_run(p.add_run()).add_break(WD_BREAK.PAGE)

    def title_page(self):
        m = self.content['metadata']
        self.para(m['university'].upper(), 'Guide Text', True, 15, 4)
        self.para(m['graduate_college'].upper(), 'Guide Text', True, 0, 35)
        self.para(m['title'], 'Title', True, 0, 40)
        self.para('A THESIS', 'Guide Text', True, 0, 8)
        self.para('SUBMITTED TO THE GRADUATE FACULTY', 'Guide Text', True, 0, 4)
        self.para('in partial fulfillment of the requirements for the', 'Guide Text', True, 0, 4)
        self.para('Degree of', 'Guide Text', True, 0, 4)
        self.para(m['degree'].upper(), 'Guide Text', True, 0, 38)
        self.para('By ' + m['author'].upper(), 'Guide Text', True, 0, 30)
        self.para(m['city'], 'Guide Text', True, 0, 4)
        self.para(str(m['year']), 'Guide Text', True)

    def committee_page(self):
        m = self.content['metadata']
        self.para(m['title'], 'Title', True, 40, 55)
        label = 'A THESIS APPROVED FOR THE' if self.args.final_approval_page else 'A THESIS SUBMITTED FOR REVIEW TO THE'
        self.para(label, 'Guide Text', True, 0, 12)
        self.para(m['college'].upper(), 'Guide Text', True, 0, 35)
        self.para('BY THE COMMITTEE CONSISTING OF', 'Guide Text', True, 0, 20)
        for name in m['committee']:
            self.para(name, 'Guide Text', True, 0, 15)
        if not self.args.final_approval_page:
            self.para('Approval of this revised draft is pending.', 'Guide Text', True, 42, 0)

    def copyright_page(self):
        m = self.content['metadata']
        # Word suppresses space-before immediately after a manual page break.
        # A spacer preserves the intended lower-page copyright placement.
        self.para('', 'Guide Text', space_after=440)
        self.para('© Copyright by ' + m['author'] + ' ' + str(m['year']), 'Guide Text', True, 0, 8)
        self.para('All Rights Reserved.', 'Guide Text', True)

    def front_heading(self, text):
        return self.para(text, 'Front Heading', center=True)

    def table_of_contents(self):
        self.front_heading('Table of Contents')
        # The major body sections exactly match the TOC field's Heading 1 source.
        entries = [('Chapter ' + str(c['number']) + '  ' + c['title'], 'chapter_' + str(c['number'])) for c in self.content['chapters']]
        entries += [('References', 'references'), ('Glossary', 'glossary')]
        entries += [(f'Appendix {chr(65+i)}  ' + a['title'], 'appendix_' + chr(65+i)) for i, a in enumerate(self.content.get('appendices', []))]
        p = self.doc.add_paragraph(style='TOC Entry')
        begin = OxmlElement('w:fldChar')
        begin.set(qn('w:fldCharType'), 'begin')
        instr = OxmlElement('w:instrText')
        instr.set(qn('xml:space'), 'preserve')
        instr.text = ' TOC \\o "1-1" \\h \\z \\u '
        separate = OxmlElement('w:fldChar')
        separate.set(qn('w:fldCharType'), 'separate')
        font_run(p.add_run())._r.append(begin)
        font_run(p.add_run())._r.append(instr)
        font_run(p.add_run())._r.append(separate)
        for title, bookmark in entries:
            p = self.para(title + '\t', 'TOC Entry')
            add_field(p, f'PAGEREF {bookmark} \\h', str(self.page_map.get(bookmark, {'front_abstract': 'iv', 'front_figures': 'vi', 'front_reader': 'vii', 'chapter_1': '1'}.get(bookmark, '?'))))
        end = OxmlElement('w:fldChar')
        end.set(qn('w:fldCharType'), 'end')
        font_run(p.add_run())._r.append(end)

    def figure_list(self):
        p = self.front_heading('List of Figures')
        self.bookmark(p, 'front_figures')
        for index, figure in enumerate(self.figures, 1):
            bookmark = 'fig_' + figure['id'].replace('-', '_')
            p = self.para(f'Figure {index}. {figure["caption"]}\t', 'TOC Entry')
            add_field(p, f'PAGEREF {bookmark} \\h', str(self.page_map.get(bookmark, '?')))
            p.paragraph_format.space_after = Pt(12)

    def reader_guide(self):
        p = self.front_heading('Reader Guide')
        self.bookmark(p, 'front_reader')
        for text in self.content['front']['reader_guide']:
            self.para(text, 'Guide Text', space_after=10)
        legend = {'columns': ['Label', 'Meaning'], 'rows': [['Blue', 'Evidence and fixed checks'], ['Purple', 'Agent choices'], ['Green', 'Supported readiness under the profile'], ['Amber', 'Missing evidence or further review'], ['Crimson', 'A known blocking problem']]}
        self.add_table(legend, caption=False, widths=[1.15, 5.35])
        self.para('Author review and committee approval of this revision remain pending. Independent label review is also pending.', 'Guide Text', space_before=8)

    def add_table(self, spec, caption=True, widths=None):
        columns = spec['columns']
        if caption:
            self.table_count += 1
            p = self.doc.add_paragraph(style='Caption')
            font_run(p.add_run('Table '))
            add_field(p, 'SEQ Table \\* ARABIC', str(self.table_count))
            font_run(p.add_run('. ' + spec['caption']))
            p.paragraph_format.keep_with_next = True
            self.bookmark(p, 'table_' + spec['id'])
        table = self.doc.add_table(rows=1, cols=len(columns))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        if widths is None:
            widths = {2: [2.1, 4.4], 3: [1.4, 2.7, 2.4], 4: [1.4, 1.3, 1.3, 2.5]}.get(len(columns), [6.5 / len(columns)] * len(columns))
        for column, width in zip(table.columns, widths):
            column.width = Inches(width)
        for cell, text, width in zip(table.rows[0].cells, columns, widths):
            cell.width = Inches(width)
            self.cell_text(cell, text, True)
        header = OxmlElement('w:tblHeader')
        table.rows[0]._tr.get_or_add_trPr().append(header)
        for row_index, row in enumerate(spec['rows']):
            cells = table.add_row().cells
            for cell, text, width in zip(cells, row, widths):
                cell.width = Inches(width)
                self.cell_text(cell, str(text), False, row_index)
        for row in table.rows:
            cant_split = OxmlElement('w:cantSplit')
            row._tr.get_or_add_trPr().append(cant_split)
        tblpr = table._tbl.tblPr
        borders = OxmlElement('w:tblBorders')
        for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            border = OxmlElement('w:' + edge)
            border.set(qn('w:val'), 'single')
            border.set(qn('w:sz'), '4')
            border.set(qn('w:color'), 'D9D9D9')
            borders.append(border)
        tblpr.append(borders)
        self.para('', 'Guide Text', space_after=3)

    def cell_text(self, cell, text, header=False, row_index=0):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        run = font_run(p.add_run(text), bold=header)
        if header:
            run.font.color.rgb = RGBColor(255, 255, 255)
        tcpr = cell._tc.get_or_add_tcPr()
        shade = OxmlElement('w:shd')
        shade.set(qn('w:fill'), '205493' if header else ('F2F5F8' if row_index % 2 == 1 else 'FFFFFF'))
        tcpr.append(shade)
        margins = OxmlElement('w:tcMar')
        for edge in ('top', 'left', 'bottom', 'right'):
            child = OxmlElement('w:' + edge)
            child.set(qn('w:w'), '95')
            child.set(qn('w:type'), 'dxa')
            margins.append(child)
        tcpr.append(margins)

    def callout(self, spec):
        accent, fill = COLORS[spec['color']]
        p = self.doc.add_paragraph(style='Guide Text')
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(8)
        font_run(p.add_run(spec['title'] + '. '), bold=True)
        font_run(p.add_run(spec['text']))
        ppr = p._p.get_or_add_pPr()
        shade = OxmlElement('w:shd')
        shade.set(qn('w:fill'), fill)
        ppr.append(shade)
        border = OxmlElement('w:pBdr')
        edge = OxmlElement('w:left')
        edge.set(qn('w:val'), 'single')
        edge.set(qn('w:sz'), '16')
        edge.set(qn('w:space'), '8')
        edge.set(qn('w:color'), accent)
        border.append(edge)
        ppr.append(border)

    def find_figure(self, figure_id):
        if figure_id in self.figure_map:
            entry = self.figure_map[figure_id]
            result = Path(entry['path'] if isinstance(entry, dict) else entry)
            return result if result.is_absolute() else Path(self.args.figure_map).resolve().parent / result
        candidates = [ROOT / 'docs' / 'assets' / 'diagrams', ROOT / 'docs' / 'assets', BASE / 'figures']
        for name in FIGURE_ALIASES.get(figure_id, [figure_id]):
            for directory in candidates:
                path = directory / (name + '.png')
                if path.exists():
                    return path
        return None

    def figure(self, spec):
        path = self.find_figure(spec['id'])
        if not path or not path.exists():
            self.missing_assets.append(spec['id'])
            if not self.args.allow_missing_figures:
                raise FileNotFoundError('Missing figure: ' + spec['id'])
        self.figure_count += 1
        if path and path.exists():
            with Image.open(path) as source:
                pixels_w, pixels_h = source.size
            entry = self.figure_map.get(spec['id'], {})
            entry = entry if isinstance(entry, dict) else {}
            diagrams = {'system-topology', 'review-sequence', 'approval-rollback', 'benchmark-split', 'migration'}
            max_height = float(entry.get('max_height_inches', spec.get('max_height_inches', 4.5 if spec['id'] in diagrams else 3.65)))
            max_width = min(6.5, float(entry.get('max_width_inches', spec.get('max_width_inches', 6.5))))
            width = min(max_width, max_height * pixels_w / pixels_h)
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Inches(0)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            shape = p.add_run().add_picture(str(path), width=Inches(width))
            shape._inline.docPr.set('descr', spec['alt'])
            shape._inline.docPr.set('title', spec['caption'])
        else:
            self.para('Diagram export pending for the rendering preview.', 'Guide Text')
        p = self.doc.add_paragraph(style='Caption')
        font_run(p.add_run('Figure '))
        add_field(p, 'SEQ Figure \\* ARABIC', str(self.figure_count))
        font_run(p.add_run('. ' + spec['caption']))
        self.bookmark(p, 'fig_' + spec['id'])

    def sections(self, sections, prefix=None):
        for index, sec in enumerate(sections, 1):
            label = f'{prefix}.{index}  ' if prefix else ''
            p = self.para(label + sec['title'], 'Heading 2')
            self.bookmark(p, f'section_{prefix}_{index}')
            for text in sec['paragraphs']:
                self.para(text)
            if sec.get('callout'):
                self.callout(sec['callout'])
            if sec.get('table'):
                self.add_table(sec['table'])
            if sec.get('figure'):
                self.figure(sec['figure'])

    def references(self):
        self.break_page()
        p = self.para('References', 'Heading 1')
        self.bookmark(p, 'references')
        # Alphabetical order, preserving source titles and verified identifiers.
        for ref in sorted(self.content['references'], key=lambda r: r['text'].casefold()):
            if self.args.print_reference_urls:
                self.para(ref['text'] + ' ' + ref['url'], 'Reference Text')
            else:
                p = self.para(ref['text'] + ' ', 'Reference Text')
                add_external_link(p, 'Source', ref['url'])

    def glossary(self):
        self.break_page()
        p = self.para('Glossary', 'Heading 1')
        self.bookmark(p, 'glossary')
        for term in self.content['glossary']:
            p = self.doc.add_paragraph(style='Glossary Text')
            font_run(p.add_run(term['term'] + '. '), bold=True)
            font_run(p.add_run(term['definition'] + ' Example: ' + term['example'] + ' Why it matters: ' + term['relevance']))
            p.paragraph_format.keep_together = True
            _, fill = COLORS[term['color']]
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), fill)
            p._p.get_or_add_pPr().append(shade)

    def build(self):
        self.title_page()
        self.break_page()
        self.committee_page()
        self.break_page()
        self.copyright_page()
        configure_section(self.doc.add_section(WD_SECTION_START.NEW_PAGE), 'lowerRoman', 4)
        p = self.front_heading('Abstract')
        self.bookmark(p, 'front_abstract')
        for text in self.content['front']['abstract']:
            self.para(text, 'Guide Text', space_after=10)
        self.para('Keywords: ' + '; '.join(self.content['front']['keywords']), 'Guide Text', space_before=8)
        self.break_page()
        self.table_of_contents()
        self.break_page()
        self.figure_list()
        self.break_page()
        self.reader_guide()
        configure_section(self.doc.add_section(WD_SECTION_START.NEW_PAGE), 'decimal', 1)
        for chapter in self.content['chapters']:
            if chapter['number'] != 1:
                self.break_page()
            p = self.para(f'Chapter {chapter["number"]}  {chapter["title"]}', 'Heading 1')
            self.bookmark(p, 'chapter_' + str(chapter['number']))
            self.sections(chapter['sections'], str(chapter['number']))
        self.references()
        self.glossary()
        for index, appendix in enumerate(self.content.get('appendices', [])):
            self.break_page()
            letter = chr(65 + index)
            p = self.para('Appendix ' + letter + '  ' + appendix['title'], 'Heading 1')
            self.bookmark(p, 'appendix_' + letter)
            self.sections(appendix['sections'], letter)
        update = self.doc.settings.element.find(qn('w:updateFields'))
        if update is None:
            update = OxmlElement('w:updateFields')
            self.doc.settings.element.append(update)
        update.set(qn('w:val'), 'true')
        output = Path(self.args.output).resolve()
        if output == Path(self.args.original).resolve():
            raise ValueError('Refusing to overwrite the confirmed completed thesis.')
        output.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(output)
        manifest = {
            'output': str(output), 'title': self.content['metadata']['title'],
            'source': str(Path(self.args.content).resolve()),
            'figure_count': self.figure_count, 'table_count': self.table_count,
            'bookmarks': self.bookmarks, 'missing_assets': self.missing_assets,
            'unresolved_placeholders_at_input': self.original_placeholders,
            'final_approval_page': self.args.final_approval_page,
            'format': {'paper': 'US Letter', 'margins_inches': 1, 'body_font': 'Times New Roman', 'body_pt': 12, 'body_line_spacing': 2, 'first_three': 'unnumbered', 'front_numbering': 'lower Roman beginning iv', 'body_numbering': 'Arabic beginning 1'},
            'validation': 'DOCX creation only. Root must render, update fields, verify page count, and inspect every page before delivery.'
        }
        output.with_suffix('.build.json').write_text(json.dumps(manifest, indent=2), encoding='utf-8')
        print(json.dumps(manifest, indent=2))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--content', default=str(BASE / 'thesis_content.json'))
    parser.add_argument('--output', default=str(ROOT / 'thesis' / '2026_Diaz_Daren_Thesis_Agent_Revision.docx'))
    parser.add_argument('--original', default=str(BASE.parent / 'Written Documentation' / 'A Conformance Harness For Verifiable Modular Open Systems Outcomes In Defense Acquisition.docx'))
    parser.add_argument('--figure-map')
    parser.add_argument('--page-map')
    parser.add_argument('--substitutions')
    parser.add_argument('--allow-placeholders', action='store_true', help='Rendering preview only. Replaces missing empirical passages with explicit pending text.')
    parser.add_argument('--allow-missing-figures', action='store_true', help='Rendering preview only. Final delivery must include verified figures.')
    parser.add_argument('--final-approval-page', action='store_true', help='Use only after actual approval of this revised thesis.')
    parser.add_argument('--print-reference-urls', action='store_true', help='Print full reference URLs. Default uses compact Source links retaining full URL targets.')
    args = parser.parse_args()
    Builder(args).build()


if __name__ == '__main__':
    main()
